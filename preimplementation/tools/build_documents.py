from pathlib import Path
import re, json, argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)
SRC = ROOT / "docs/TriX_Design_and_Implementation_Standard_v0_2.md"
DEST = OUT / "TriX_Design_and_Implementation_Standard_v0_2.docx"
PAGE_MAP_PATH = ROOT / "validation/page_map.json"
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.25)
sec.footer_distance = Inches(0.30)

FONT = "Nimbus Sans"
MONO = "DejaVu Sans Mono"
BLACK = RGBColor.from_string("000000")
styles = doc.styles
for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Heading 3","Caption","Header","Footer"]:
    styles[name].font.name = FONT
    styles[name].font.color.rgb = BLACK
    rp=styles[name].element.get_or_add_rPr()
    rf=rp.find(qn("w:rFonts"))
    if rf is None:
        rf=OxmlElement("w:rFonts");rp.append(rf)
    for field in ["ascii","hAnsi","eastAsia","cs"]: rf.set(qn("w:"+field),FONT)
    clr=rp.find(qn("w:color"))
    if clr is not None:
        for a in ["themeColor","themeTint","themeShade"]:
            clr.attrib.pop(qn("w:"+a),None)

normal=styles["Normal"]
normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.widow_control=True
styles["Title"].font.size=Pt(32)
styles["Title"].font.bold=True
styles["Title"].paragraph_format.space_after=Pt(12)
styles["Subtitle"].font.size=Pt(21)
styles["Subtitle"].paragraph_format.space_after=Pt(14)
for name,size,before,after in [("Heading 1",16,16,9),("Heading 2",12.5,12,6),("Heading 3",11.5,10,5)]:
    st=styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(before)
    st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.keep_together=True
    st.paragraph_format.widow_control=True

def newstyle(name,base="Normal"):
    if name not in styles: styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st=styles[name];st.base_style=styles[base];st.font.name=FONT;st.font.color.rgb=BLACK
    return st

reqstyle=newstyle("Requirement")
reqstyle.paragraph_format.space_after=Pt(7)
reqstyle.paragraph_format.keep_together=True
casestyle=newstyle("Acceptance")
casestyle.paragraph_format.space_after=Pt(7)
casestyle.paragraph_format.keep_together=True
codes=newstyle("Code")
codes.font.name=MONO;codes.font.size=Pt(9)
codes.paragraph_format.line_spacing=1.05
codes.paragraph_format.space_after=Pt(0)
codes.paragraph_format.keep_together=True
codes.paragraph_format.left_indent=Inches(0.15)
meta=newstyle("Metadata");meta.font.size=Pt(10)
meta.paragraph_format.space_after=Pt(18)
tocstyle=newstyle("Contents Entry")
tocstyle.font.size=Pt(11)
tocstyle.paragraph_format.space_after=Pt(5)
tocstyle.paragraph_format.keep_together=True
tocstyle.paragraph_format.tab_stops.add_tab_stop(Inches(6.76), WD_ALIGN_PARAGRAPH.RIGHT)
srcstyle=newstyle("Source")
srcstyle.font.size=Pt(10)
srcstyle.paragraph_format.line_spacing=1.08
srcstyle.paragraph_format.space_after=Pt(2)
srcstyle.paragraph_format.keep_with_next=True
urlstyle=newstyle("Source URL")
urlstyle.font.size=Pt(9)
urlstyle.paragraph_format.space_after=Pt(9)
urlstyle.paragraph_format.line_spacing=1.05
urlstyle.paragraph_format.keep_together=True
tabstyle=newstyle("Table text")
tabstyle.font.size=Pt(9.5)
tabstyle.paragraph_format.line_spacing=1.08
tabstyle.paragraph_format.space_before=Pt(0)
tabstyle.paragraph_format.space_after=Pt(0)
tabstyle.paragraph_format.widow_control=True

# Explicit document language and predictable typography.
defaults=styles.element.find(qn("w:docDefaults"))
if defaults is not None:
    rpd=defaults.find(qn("w:rPrDefault"))
    if rpd is not None:
        rp=rpd.find(qn("w:rPr"))
        if rp is not None:
            lang=rp.find(qn("w:lang"))
            if lang is None: lang=OxmlElement("w:lang");rp.append(lang)
            lang.set(qn("w:val"),"en-GB")
doc.core_properties.title="TriX Design and Implementation Standard"
doc.core_properties.subject="TriX design requirements, execution contracts and acceptance gates"
doc.core_properties.author="TriX Project"
doc.core_properties.keywords="TriX, language, requirements, clusters, platforms, conformance"
doc.core_properties.comments=""
doc.core_properties.version="0.2"
settings=doc.settings.element
upd=OxmlElement("w:updateFields");upd.set(qn("w:val"),"true");settings.append(upd)

bookmark_id=0
def bookmark(p,name):
    global bookmark_id
    bookmark_id+=1
    start=OxmlElement("w:bookmarkStart");start.set(qn("w:id"),str(bookmark_id));start.set(qn("w:name"),name)
    end=OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),str(bookmark_id))
    p._p.insert(0,start);p._p.append(end)

def hyperlink(p,text,anchor=None,url=None,size=None,color="000000",bold=False):
    h=OxmlElement("w:hyperlink")
    if anchor: h.set(qn("w:anchor"),anchor)
    if url: h.set(qn("r:id"),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement("w:r");rp=OxmlElement("w:rPr")
    rf=OxmlElement("w:rFonts");rf.set(qn("w:ascii"),FONT);rf.set(qn("w:hAnsi"),FONT);rp.append(rf)
    c=OxmlElement("w:color");c.set(qn("w:val"),color);rp.append(c)
    if size:
        sz=OxmlElement("w:sz");sz.set(qn("w:val"),str(int(size*2)));rp.append(sz)
    if bold:rp.append(OxmlElement("w:b"))
    r.append(rp);t=OxmlElement("w:t");t.text=text;t.set(qn("xml:space"),"preserve");r.append(t);h.append(r);p._p.append(h)

def add_inline(p,text):
    # Bold labels and hyperlinks to bibliography entries.
    for item in re.split(r"(\*\*.*?\*\*|\[F\d{2}\])",text):
        if not item:continue
        if item.startswith("**") and item.endswith("**"):
            p.add_run(item[2:-2]).bold=True
        elif re.fullmatch(r"\[F\d{2}\]",item):
            hyperlink(p,item,anchor=item[1:-1])
        else:p.add_run(item)

def add_para(text,style=None):
    p=doc.add_paragraph(style=style)
    add_inline(p,text)
    return p

def cell_shading(cell,fill):
    pr=cell._tc.get_or_add_tcPr()
    sh=pr.find(qn("w:shd"))
    if sh is None:sh=OxmlElement("w:shd");pr.append(sh)
    sh.set(qn("w:fill"),fill);sh.set(qn("w:val"),"clear")

def add_table(rows):
    n=len(rows[0])
    table=doc.add_table(rows=1,cols=n)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths = [1.57,5.29] if n==2 else ([1.72,2.28,2.86] if n==3 else [6.86/n]*n)
    for col,w in zip(table.columns,widths):col.width=Inches(w)
    pr=table._tbl.tblPr
    borders=OxmlElement("w:tblBorders")
    for edge in ["top","left","bottom","right","insideH","insideV"]:
        e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
    pr.append(borders)
    margins=OxmlElement("w:tblCellMar")
    for edge,value in [("top",95),("bottom",95),("left",105),("right",105)]:
        e=OxmlElement("w:"+edge);e.set(qn("w:w"),str(value));e.set(qn("w:type"),"dxa");margins.append(e)
    pr.append(margins)
    for i,values in enumerate(rows):
        row=table.rows[0] if i==0 else table.add_row()
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement("w:cantSplit");trpr.append(cant)
        if i==0:
            repeat=OxmlElement("w:tblHeader");repeat.set(qn("w:val"),"true");trpr.append(repeat)
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.style=styles["Table text"]
            add_inline(p,value)
            cell_shading(cell,"26313D" if i==0 else ("F2F5F7" if i%2==0 else "FFFFFF"))
            if i==0:
                for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
            if n==2 and j==0 and len(value)<14:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    spacer=doc.add_paragraph()
    spacer.paragraph_format.space_after=Pt(1)
    spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(3)
    spacer.add_run().font.size=Pt(3)
    return table

# Footer is functional document navigation.
footer=sec.footer.paragraphs[0]
footer.style=styles["Footer"]
styles["Footer"].paragraph_format.tab_stops.clear_all()
footer.paragraph_format.tab_stops.clear_all()
footer.paragraph_format.space_after=Pt(0)
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.86),WD_ALIGN_PARAGRAPH.RIGHT)
r=footer.add_run("TriX | Design and Implementation Standard | v0.2");r.font.size=Pt(9)
footer.add_run("\t")
for text,isfield in [("Page ",False),("PAGE",True),(" of ",False),("NUMPAGES",True)]:
    if not isfield:
        r=footer.add_run(text);r.font.size=Pt(9)
    else:
        r=footer.add_run();r.font.size=Pt(9)
        b=OxmlElement("w:fldChar");b.set(qn("w:fldCharType"),"begin");r._r.append(b)
        ins=OxmlElement("w:instrText");ins.set(qn("xml:space"),"preserve");ins.text=" "+text+" ";r._r.append(ins)
        sep=OxmlElement("w:fldChar");sep.set(qn("w:fldCharType"),"separate");r._r.append(sep)
        t=OxmlElement("w:t");t.text="1";r._r.append(t)
        e=OxmlElement("w:fldChar");e.set(qn("w:fldCharType"),"end");r._r.append(e)

source=SRC.read_text(encoding="utf-8")
lines=source.splitlines()
chapter_start=next(i for i,l in enumerate(lines) if l.startswith("# 1 "))
front="\n".join(lines[:chapter_start]).strip().split("\n\n")
doc.add_paragraph("TriX",style="Title")
doc.add_paragraph("Design and Implementation Standard",style="Subtitle")
doc.add_paragraph("Version 0.2 | 22 October 2026\nProposed normative baseline",style="Metadata")
for part in front[2:]:
    if part.startswith("Version"):continue
    if part.startswith("#"):continue
    add_para(part)
doc.add_heading("Using this standard",level=2)
add_para("Adopt a revision and select the applicable profiles. Link every mandatory requirement to implementation and validation evidence. Chapter 19 sets out delivery gates and decisions that remain open.")
add_para("The companion pack provides implementation guidance, structured contracts, examples and registers. Its examples are illustrative; schema checks do not certify a TriX implementation.")
reqs=re.findall(r"\*\*([A-Z]{3}-\d{3})",source)
cases=re.findall(r"^\*\*((?:CA|RT|QT)-\d{2}|(?:LSP|SIF)-AT\d{2})",source,re.M)
chapter_count=len(re.findall(r"^# \d+ ",source,re.M))
reference_count=len(re.findall(r"^\*\*F\d{2}",source,re.M))
add_para(f"Contents: {len(reqs)} identified requirements, {len(cases)} acceptance cases, {chapter_count} chapters and {reference_count} technical references.",style="Metadata")
doc.add_page_break()

doc.add_heading("Contents",level=1)
headings=[l[2:] for l in lines if re.match(r"# \d+ ",l)]
page_map=json.loads(PAGE_MAP_PATH.read_text()) if PAGE_MAP_PATH.exists() else {}
for title in headings:
    num=title.split(" ",1)[0]
    p=doc.add_paragraph(style="Contents Entry")
    hyperlink(p,title,anchor="chapter_"+num)
    p.add_run("\t"+str(page_map.get(num,"")))
add_para("Contents entries and bracketed source references link to their destinations in the editable edition.",style="Metadata")
doc.add_page_break()

in_code=False
code_lines=[]
table_lines=[]
paragraph_lines=[]
current_chapter=""
def flush_paragraph():
    global paragraph_lines
    if not paragraph_lines:return
    text=" ".join(paragraph_lines);paragraph_lines=[]
    if re.match(r"^\*\*[A-Z]{3}-\d{3}",text):
        p=add_para(text,"Requirement")
        bookmark(p,re.match(r"^\*\*([A-Z]{3}-\d{3})",text).group(1).replace("-","_"))
    elif re.match(r"^\*\*((?:CA|RT|QT)-\d{2}|(?:LSP|SIF)-AT\d{2})",text):
        p=add_para(text,"Acceptance")
        bookmark(p,re.match(r"^\*\*((?:CA|RT|QT)-\d{2}|(?:LSP|SIF)-AT\d{2})",text).group(1).replace("-","_"))
    elif re.match(r"^\*\*F\d{2}",text):
        p=add_para(text,"Source")
        bookmark(p,re.match(r"^\*\*(F\d{2})",text).group(1))
    elif text.startswith("https://") or re.fullmatch(r"\[Source\]\((https://[^)]+)\)",text):
        source_link=re.fullmatch(r"\[Source\]\((https://[^)]+)\)",text)
        url=source_link.group(1) if source_link else text
        p=doc.add_paragraph(style="Source URL")
        hyperlink(p,url,url=url,size=9,color="174A70")
    else:add_para(text)

def flush_table():
    global table_lines
    if not table_lines:return
    rows=[]
    for l in table_lines:
        vals=[x.strip() for x in l.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?",x) for x in vals):continue
        rows.append(vals)
    add_table(rows);table_lines=[]

for line in lines[chapter_start:]:
    if line.startswith("~~~"):
        flush_paragraph();flush_table()
        if in_code:
            for idx,cl in enumerate(code_lines):
                p=doc.add_paragraph(style="Code");p.add_run(cl)
                p.paragraph_format.keep_with_next=idx<len(code_lines)-1
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            code_lines=[]
        in_code=not in_code
        continue
    if in_code:code_lines.append(line);continue
    if line.startswith("|"):
        flush_paragraph();table_lines.append(line);continue
    if table_lines:flush_table()
    if line.startswith("# "):
        flush_paragraph()
        text=line[2:]
        p=doc.add_heading(text,level=1)
        bookmark(p,"chapter_"+text.split(" ",1)[0])
        current_chapter=text
        continue
    if re.match(r"^(?:[-*] |\d+\. )",line):
        flush_paragraph()
        p=doc.add_paragraph(style="List Bullet" if line.startswith(("- ","* ")) else "Normal")
        add_inline(p,line[2:] if line.startswith(("- ","* ")) else line)
        continue
    if line.startswith("## "):
        flush_paragraph();doc.add_heading(line[3:],level=2);continue
    if not line.strip():
        flush_paragraph();continue
    if line.startswith("https://"):
        flush_paragraph();paragraph_lines=[line];flush_paragraph();continue
    paragraph_lines.append(line)
flush_paragraph();flush_table()

# No title or paragraph borders are part of this design.
for p in doc.paragraphs:
    ppr=p._p.pPr
    if ppr is not None:
        borders=ppr.find(qn("w:pBdr"))
        if borders is not None:ppr.remove(borders)
for name in ["Title","Subtitle","Heading 1","Heading 2","Heading 3"]:
    ppr=styles[name].element.find(qn("w:pPr"))
    if ppr is not None:
        b=ppr.find(qn("w:pBdr"))
        if b is not None:ppr.remove(b)
doc.save(DEST)
print(json.dumps({"path":str(DEST),"requirements":len(reqs),"scenarios":len(cases),"chapters":len(headings),"bytes":DEST.stat().st_size},ensure_ascii=False))

